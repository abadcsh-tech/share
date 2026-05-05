#!/usr/bin/env python3
"""Build devotional_v2.pptx — Anthropic claude.com design + Korean fonts (NanumMyeongjo ExtraBold + SUIT Bold).

Critical: sets latin/ea/cs typeface so Korean characters render correctly in PowerPoint.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ============================================================
# Tokens — Anthropic claude.com palette
# ============================================================
C = {
    "primary":      RGBColor(0xCC, 0x78, 0x5C),   # coral
    "canvas":       RGBColor(0xFA, 0xF9, 0xF5),   # cream
    "card":         RGBColor(0xEF, 0xE9, 0xDE),   # cream card
    "card_strong":  RGBColor(0xE8, 0xE0, 0xD2),
    "dark":         RGBColor(0x18, 0x17, 0x15),   # navy/black
    "dark_elev":    RGBColor(0x25, 0x23, 0x20),
    "ink":          RGBColor(0x14, 0x14, 0x13),
    "body":         RGBColor(0x3D, 0x3D, 0x3A),
    "muted":        RGBColor(0x6C, 0x6A, 0x64),
    "on_primary":   RGBColor(0xFF, 0xFF, 0xFF),
    "on_dark":      RGBColor(0xFA, 0xF9, 0xF5),
    "on_dark_soft": RGBColor(0xA0, 0x9D, 0x96),
    "hairline":     RGBColor(0xE6, 0xDF, 0xD8),
}

FONT_DISPLAY = "NanumMyeongjo"   # 나눔명조 ExtraBold (style)
FONT_SANS = "SUIT"               # SUIT Bold

# ============================================================
# Korean font handling — must set latin AND eastAsia AND cs
# This is the trap the user warned about: setting only run.font.name
# leaves Korean text to fall back to default East Asian font.
# ============================================================
def set_korean_font(run, font_name, bold=False):
    rPr = run._r.get_or_add_rPr()
    if bold:
        rPr.set("b", "1")
    # Remove existing typeface elements
    for tag in ["a:latin", "a:ea", "a:cs"]:
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    # Add latin
    latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font_name)
    # Add eastAsia (Korean falls under this)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font_name)
    # Add complex script
    cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", font_name)

def add_text_box(slide, x, y, w, h, text, *, font=FONT_SANS, size=18, bold=False,
                 color=None, align="left", valign="top", line_spacing=1.4, char_spacing=0):
    if color is None:
        color = C["ink"]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    if valign == "middle":
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif valign == "bottom":
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align == "center":
            p.alignment = PP_ALIGN.CENTER
        elif align == "right":
            p.alignment = PP_ALIGN.RIGHT
        else:
            p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        # Apply Korean-aware font
        set_korean_font(run, font, bold=bold)
        if char_spacing:
            run._r.get_or_add_rPr().set("spc", str(char_spacing))
    return box

def add_rect(slide, x, y, w, h, fill, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape

def add_rounded_rect(slide, x, y, w, h, fill, line_color=None, corner_pct=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    # Set corner radius via adjustment
    try:
        shape.adjustments[0] = corner_pct
    except Exception:
        pass
    return shape

# ============================================================
# Layout — 16:9 at 13.333 × 7.5 inch
# ============================================================
W, H = 13.333, 7.5
PAD = 0.85
CW = W - 2 * PAD

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

BLANK = prs.slide_layouts[6]
TOTAL = 16

# ============================================================
# Helper: bg fill
# ============================================================
def bg(slide, color):
    bg_shape = add_rect(slide, 0, 0, W, H, color)
    # Move to back
    spTree = slide.shapes._spTree
    spTree.remove(bg_shape._element)
    spTree.insert(2, bg_shape._element)

def progress_bar(slide, n, total, dark=False):
    track = RGBColor(0x33, 0x33, 0x33) if dark else RGBColor(0xEE, 0xEC, 0xE6)
    add_rect(slide, 0, 0, W, 0.05, track)
    add_rect(slide, 0, 0, W * n / total, 0.05, C["primary"])
    cap_color = C["on_dark_soft"] if dark else C["muted"]
    add_text_box(slide, PAD, H - 0.5, 6, 0.3,
                 f"한 가족이 되신 예수님 · 5일 묵상",
                 font=FONT_SANS, size=10, bold=True, color=cap_color, char_spacing=300)
    add_text_box(slide, W - 1.6, H - 0.5, 1.4, 0.3,
                 f"{n:02d} / {total:02d}",
                 font=FONT_SANS, size=10, bold=True, color=cap_color, align="right", char_spacing=300)

def label(slide, text, dark=False, with_spike=True, y=PAD - 0.05):
    color = C["on_dark"] if dark else C["muted"]
    spike_color = C["primary"] if not dark else C["primary"]
    if with_spike:
        # Spike + text
        add_text_box(slide, PAD, y, 0.4, 0.4, "✸", font=FONT_SANS, size=16, bold=True, color=spike_color)
        add_text_box(slide, PAD + 0.4, y + 0.05, CW - 0.4, 0.4,
                     text, font=FONT_SANS, size=18, bold=True, color=color, char_spacing=400)
    else:
        add_text_box(slide, PAD, y + 0.05, CW, 0.4,
                     text, font=FONT_SANS, size=18, bold=True, color=color, char_spacing=400)

def hero_title(slide, line1, line2, dark=False, y=1.5):
    color = C["on_dark"] if dark else C["ink"]
    add_text_box(slide, PAD, y, CW, 2.6,
                 f"{line1}\n{line2}",
                 font=FONT_DISPLAY, size=72, bold=True, color=color,
                 line_spacing=1.06)

def subhead(slide, text, dark=False, y=4.5):
    color = C["on_dark"] if dark else C["body"]
    add_text_box(slide, PAD, y, CW * 0.85, 0.9,
                 text, font=FONT_SANS, size=28, bold=True, color=color,
                 line_spacing=1.4)

def scripture_dark_card(slide, x, y, w, h, quote, cite):
    add_rounded_rect(slide, x, y, w, h, C["dark"], corner_pct=0.05)
    add_text_box(slide, x + 0.4, y + 0.4, w - 0.8, h * 0.6,
                 f'"{quote}"', font=FONT_DISPLAY, size=36, bold=True,
                 color=C["on_dark"], line_spacing=1.32)
    add_text_box(slide, x + 0.4, y + h - 0.55, w - 0.8, 0.4,
                 cite, font=FONT_SANS, size=21, bold=True,
                 color=C["on_dark_soft"], char_spacing=400)

def scripture_open(slide, x, y, w, quote, cite, dark=False):
    """Scripture with left coral border (no card)."""
    add_rect(slide, x, y, 0.06, 1.6, C["primary"])
    color_q = C["on_dark"] if dark else C["ink"]
    color_c = C["on_dark_soft"] if dark else C["muted"]
    add_text_box(slide, x + 0.22, y, w - 0.22, 1.2,
                 f'"{quote}"', font=FONT_DISPLAY, size=42, bold=True,
                 color=color_q, line_spacing=1.32)
    add_text_box(slide, x + 0.22, y + 1.25, w - 0.22, 0.4,
                 cite, font=FONT_SANS, size=22, bold=True,
                 color=color_c, char_spacing=400)

def display_md(slide, line1, line2, dark=False, y=1.5):
    color = C["on_dark"] if dark else C["ink"]
    add_text_box(slide, PAD, y, CW * 0.9, 2.0,
                 f"{line1}\n{line2}",
                 font=FONT_DISPLAY, size=52, bold=True, color=color,
                 line_spacing=1.2)

def body_text(slide, text, dark=False, y=4.0, h=2.5):
    color = C["on_dark_soft"] if dark else C["body"]
    add_text_box(slide, PAD, y, CW * 0.85, h,
                 text, font=FONT_SANS, size=24, bold=True, color=color,
                 line_spacing=1.55)

def question_card(slide, x, y, w, h, num, text, on_color="card"):
    fill_map = {
        "card": (RGBColor(0xFF, 0xFF, 0xFF), C["hairline"]),
        "dark": (C["dark_elev"], None),
        "coral": (RGBColor(0xFF, 0xFF, 0xFF), None),  # white-ish on coral
    }
    fill, border = fill_map.get(on_color, (RGBColor(0xFF, 0xFF, 0xFF), C["hairline"]))
    if on_color == "coral":
        # Use semi-translucent appearance via lighter coral overlay
        fill = RGBColor(0xE6, 0x8E, 0x71)  # softer coral
    rect = add_rounded_rect(slide, x, y, w, h, fill, line_color=border, corner_pct=0.06)
    text_color = C["on_dark"] if on_color == "dark" else (C["on_primary"] if on_color == "coral" else C["ink"])
    num_color = C["primary"] if on_color != "coral" else C["on_primary"]
    add_text_box(slide, x + 0.3, y + 0.18, 0.5, 0.5,
                 num, font="JetBrains Mono", size=14, bold=True, color=num_color)
    add_text_box(slide, x + 0.95, y + 0.16, w - 1.1, h - 0.3,
                 text, font=FONT_SANS, size=20, bold=True, color=text_color,
                 line_spacing=1.5, valign="middle")

def prayer_text(slide, text, dark=False, y=5.5):
    color = C["on_dark"] if dark else C["ink"]
    add_text_box(slide, PAD, y, CW * 0.85, 1.5,
                 text, font=FONT_DISPLAY, size=30, bold=True, color=color,
                 line_spacing=1.45)

def amen(slide, dark=False, y=H - 0.95):
    color = C["on_dark_soft"] if dark else C["muted"]
    add_text_box(slide, PAD, y, 3, 0.3,
                 "— 아멘.", font=FONT_SANS, size=20, bold=True, color=color, char_spacing=400)

# ============================================================
# Slide 1: Series intro
# ============================================================
s = prs.slides.add_slide(BLANK)
bg(s, C["canvas"])
label(s, "사도신경 강해 10 · 죽으시고, 장사된 지")
add_text_box(s, PAD, 1.4, CW, 2.0,
             "한 가족이 되신 예수님,\n다섯 날의 묵상.",
             font=FONT_DISPLAY, size=60, bold=True, color=C["ink"], line_spacing=1.1)
add_text_box(s, PAD, 3.6, CW * 0.7, 0.6,
             "매일 한 편씩 — 본문 한 구절, 핵심 한 문장, 짧은 묵상글과 기도.",
             font=FONT_SANS, size=22, bold=True, color=C["body"])

tile_w = (CW - 0.6) / 5
tile_h = 2.6
tile_y = 4.5
tiles = [
    ("DAY 01", "고후 8:9", "한 가족이 되신\n예수님.", False),
    ("DAY 02", "고후 8:9", "가난과 부요의\n교환.", False),
    ("DAY 03", "갈 3:13-14", "저주를 짊어지신\n예수님.", False),
    ("DAY 04", "고후 5:21", "의의 전가, 새로운\n신분.", True),  # featured
    ("DAY 05", "롬 8:38-39", "죽음조차 끊지\n못하는 사랑.", False),
]
for i, (num, scrip, ttl, featured) in enumerate(tiles):
    x = PAD + i * (tile_w + 0.15)
    fill_color = C["dark"] if featured else C["canvas"]
    border = None if featured else C["hairline"]
    add_rounded_rect(s, x, tile_y, tile_w, tile_h, fill_color, line_color=border, corner_pct=0.06)
    text_main = C["on_dark"] if featured else C["ink"]
    text_meta = C["on_dark_soft"] if featured else C["muted"]
    add_text_box(s, x + 0.3, tile_y + 0.3, tile_w - 0.6, 0.3,
                 num, font=FONT_SANS, size=14, bold=True, color=C["primary"], char_spacing=600)
    add_text_box(s, x + 0.3, tile_y + 0.65, tile_w - 0.6, 0.3,
                 scrip, font=FONT_SANS, size=13, bold=True, color=text_meta, char_spacing=500)
    add_text_box(s, x + 0.3, tile_y + tile_h - 1.4, tile_w - 0.6, 1.2,
                 ttl, font=FONT_DISPLAY, size=22, bold=True, color=text_main,
                 line_spacing=1.18, valign="bottom")

progress_bar(s, 1, TOTAL)

# ============================================================
# Day-builder
# ============================================================
def build_day(num, label_txt, ttl1, ttl2, sub, scrip, scrip_cite, key1, key2,
              med_p1, med_p2, app_h, q1, q2, prayer, hero_bg="canvas",
              med_bg="card", app_bg="dark"):
    base_idx = 1 + (num - 1) * 3 + 1  # 2, 5, 8, 11, 14

    # ---- Hero slide ----
    s = prs.slides.add_slide(BLANK)
    is_dark_h = hero_bg == "dark"
    is_coral_h = hero_bg == "coral"
    bg_color = {"canvas": C["canvas"], "card": C["card"], "dark": C["dark"], "coral": C["primary"]}[hero_bg]
    bg(s, bg_color)
    label(s, f"DAY 0{num} · {label_txt}", dark=(is_dark_h or is_coral_h))
    hero_title(s, ttl1, ttl2, dark=(is_dark_h or is_coral_h))
    subhead(s, sub, dark=(is_dark_h or is_coral_h))
    if hero_bg in ("canvas", "card"):
        scripture_dark_card(s, PAD, 5.5, CW * 0.85, 1.6, scrip, scrip_cite)
    elif hero_bg == "dark":
        scripture_open(s, PAD, 5.5, CW * 0.85, scrip, scrip_cite, dark=True)
    else:  # coral
        scripture_open(s, PAD, 5.5, CW * 0.85, scrip, scrip_cite, dark=True)
    progress_bar(s, base_idx, TOTAL, dark=(is_dark_h or is_coral_h))

    # ---- Meditation slide ----
    s = prs.slides.add_slide(BLANK)
    is_dark_m = med_bg == "dark"
    bg_color = {"canvas": C["canvas"], "card": C["card"], "dark": C["dark"]}[med_bg]
    bg(s, bg_color)
    label(s, f"DAY 0{num} · 묵상", dark=is_dark_m, with_spike=False)
    display_md(s, key1, key2, dark=is_dark_m)
    body_text(s, med_p1, dark=is_dark_m, y=4.0, h=1.3)
    body_text(s, med_p2, dark=is_dark_m, y=5.4, h=1.6)
    progress_bar(s, base_idx + 1, TOTAL, dark=is_dark_m)

    # ---- Application + prayer slide ----
    s = prs.slides.add_slide(BLANK)
    is_dark_a = app_bg == "dark"
    is_coral_a = app_bg == "coral"
    bg_color = {"canvas": C["canvas"], "card": C["card"], "dark": C["dark"], "coral": C["primary"]}[app_bg]
    bg(s, bg_color)
    label(s, f"DAY 0{num} · 적용 + 기도", dark=(is_dark_a or is_coral_a))
    add_text_box(s, PAD, 1.5, CW * 0.9, 1.6, app_h,
                 font=FONT_DISPLAY, size=42, bold=True,
                 color=(C["on_dark"] if (is_dark_a or is_coral_a) else C["ink"]),
                 line_spacing=1.25)
    qcolor = "dark" if is_dark_a else ("coral" if is_coral_a else "card")
    qW = CW * 0.85
    question_card(s, PAD, 3.2, qW, 0.95, "01", q1, on_color=qcolor)
    question_card(s, PAD, 4.3, qW, 0.95, "02", q2, on_color=qcolor)
    prayer_text(s, prayer, dark=(is_dark_a or is_coral_a), y=5.6)
    amen(s, dark=(is_dark_a or is_coral_a))
    progress_bar(s, base_idx + 2, TOTAL, dark=(is_dark_a or is_coral_a))

# ============================================================
# Days 1-5
# ============================================================
build_day(1, "한 가족 되심",
    "한 가족이", "되신 예수님.",
    "예수님은 우리를 도와주신 게 아니라, 우리와 한 가족이 되셨습니다.",
    "그리스도께서는 부요하나, 여러분을 위해서 가난하게 되셨습니다.",
    "고린도후서 8:9",
    "가족이 된다는 건,", "모든 것을 함께 나눈다는 뜻입니다.",
    "부모님 카드 속 돈은 부모님이 번 것이지만, 가족이라는 이유 하나로 내가 쓸 수 있습니다. 가족이 된다는 건 그런 겁니다.",
    "예수님이 우리에게 오신 방식이 그렇습니다. 우리의 모든 것 — 가난도, 저주도, 죄도 — 자기 것으로 가져가시고, 그분의 부요와 의로움을 우리에게 선물로 주셨습니다.",
    "한 가지만 골라, 오늘 마음에 두기.",
    "오늘 내가 예수님께 \"가족이라서\" 가져가시도록 맡길 수 있는 한 가지는?",
    "\"가족이 되어 주셨다\"는 표현을 들었을 때, 가장 먼저 떠오르는 장면은?",
    "예수님, 저를 돕는 분이 아니라 한 가족이 되어 주셔서\n감사합니다. 오늘 제 짐을 혼자 지지 않게 하소서.",
    hero_bg="canvas", med_bg="card", app_bg="dark")

build_day(2, "부요와 가난의 교환",
    "가난과", "부요의 교환.",
    "그분이 가난해지셨기에, 나는 부요해졌습니다.",
    "그의 가난으로 여러분을 부요하게 하시려는 것입니다.",
    "고린도후서 8:9",
    "도와주신 게 아니라,", "내 자리에 직접 내려오셨습니다.",
    "예수님은 후원자가 아니라 우리 자리에 직접 내려오신 분입니다. 그분이 가난하셨기에 우리는 부요해졌습니다.",
    "가난은 단지 돈이 없는 것만 아닙니다. 사랑받지 못함, 인정받지 못함, 안전감의 부재 — 이 모든 가난을 그분이 먼저 지나가셨습니다.",
    "내 안의 \"가난\"을 한 단어로 적어보기.",
    "내가 지금 느끼는 \"가난\"은? 그것이 예수님이 이미 지나가신 자리임을 떠올려 보세요.",
    "오늘 누군가에게 \"내가 가진 부요\"를 작은 형태로라도 나눌 수 있는 한 가지는?",
    "예수님, 저를 위해 가난해지신 그 사랑을 묵상합니다.\n오늘도 누군가에게 작은 부요를 흘려보내게 하소서.",
    hero_bg="card", med_bg="canvas", app_bg="coral")

build_day(3, "저주의 자리",
    "저주를", "짊어지신 예수님.",
    "가장 큰 저주의 자리에 그분이 가셨기에, 가장 큰 축복이 내게 흘러옵니다.",
    "그리스도께서 우리를 위하여 저주를 받은 사람이 되심으로써, 우리를 율법의 저주에서 속량해 주셨습니다.",
    "갈라디아서 3:13",
    "예수님은 두려운 자리에", "먼저 가신 분입니다.",
    "십자가는 단순한 처형 도구가 아니었습니다. 당시 사람들에게 그것은 \"하나님께 버림받았다\"는 가장 큰 모욕이자 저주였습니다.",
    "그 자리에 예수님이 가셨습니다. 자기 잘못이 하나도 없으셨는데도 — 우리가 받아야 할 저주를 자기 등에 지셨습니다.",
    "내 안의 \"버림받음\"의 그림자를 마주하기.",
    "내가 두려워하는 \"버림받음\"의 형태는? 그 자리에도 예수님이 먼저 가셨다는 사실이 위로가 됩니까?",
    "오늘 내가 누군가의 \"저주\" — 외로움이나 수치 — 를 함께 짊어질 수 있다면 무엇입니까?",
    "예수님, 가장 깊은 어둠의 자리에 먼저 가셔서 감사합니다.\n그 발자국을 따라 한 걸음 내딛게 하소서.",
    hero_bg="dark", med_bg="canvas", app_bg="card")

build_day(4, "새 신분",
    "의의 전가,", "새로운 신분.",
    "내가 의로워진 게 아니라, 그분의 의로움이 내게 입혀졌습니다.",
    "그것은 우리가 그리스도 안에서 하나님의 의가 되게 하시려는 것입니다.",
    "고린도후서 5:21",
    "하나님은 내가 한 일이 아니라,", "예수님이 하신 일을 보십니다.",
    "\"내가 노력해서 깨끗해지면 하나님이 받아주시겠지\" — 우리는 자주 이렇게 생각합니다. 그러나 본문은 정반대로 말합니다.",
    "죄가 전혀 없으신 분이 우리 자리에 죄인으로 서셨고, 그래서 죄 많은 우리가 그분 안에서 의로운 사람이 되었습니다. 이것은 거래가 아니라 가족의 교환입니다.",
    "\"내 의로움\"을 내려놓는 한 순간.",
    "\"내 의로움\"이 아니라 \"그분의 의로움\" 안에서 산다면, 무엇이 가장 가벼워질까요?",
    "누군가를 그의 행동이 아니라 \"예수님 안의 신분\"으로 봐주는 한 번의 시선을 가져 보세요.",
    "예수님, 제가 한 일이 아니라 당신이 하신 일을 의지하게 하소서.\n당신 안의 새 신분으로 살게 하소서.",
    hero_bg="canvas", med_bg="card", app_bg="dark")

build_day(5, "끊어지지 않는 사랑",
    "죽음조차", "끊지 못하는 사랑.",
    "그분이 죽음의 자리까지 가셨기에, 그곳도 더는 우리를 끊지 못합니다.",
    "그 어떤 피조물도, 우리를 우리 주 예수 그리스도 안에 있는 하나님의 사랑에서 끊을 수 없습니다.",
    "로마서 8:38-39",
    "예수님이 먼저 가신 자리는,", "더 이상 우리를 끊지 못합니다.",
    "예수님의 죽음은 연극이 아니었습니다. 진짜로 심장이 멈추고, 무덤에 묻히신 사건입니다. 우리와 완전히 하나가 되시기 위해서였습니다.",
    "외로움, 실패, 불안, 두려움 — 우리가 가는 모든 자리에 예수님은 먼저 가셨고, 거기서 우리에게 손을 내미십니다.",
    "한 주의 묵상에서 한 가지를 챙겨가기.",
    "이번 한 주 묵상에서 \"예수님이 먼저 가신 자리\"라는 표현이 가장 깊이 와닿은 순간은?",
    "다음 한 주, 어떤 한 가지 영역에서 이 진리를 살아내고 싶습니까?",
    "예수님, 죽음의 자리까지 가신 그 사랑이\n제 매일의 자리가 되게 하소서.",
    hero_bg="coral", med_bg="canvas", app_bg="dark")

# Save
out_path = "/Users/csh/Projects/claude-code-md-explainer/devotion-10/devotional_v2.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slide count: {len(prs.slides)}")
